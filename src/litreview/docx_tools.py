"""Low-level python-docx helpers for filling a template without damaging it.

The college template carries all of its formatting as direct run properties on
``Normal`` paragraphs rather than as named heading styles, so new content has to
be inserted *next to* an existing paragraph and inherit its properties. Every
helper here works that way: find an anchor, clone its formatting, insert after it.

Page numbers cannot be computed by python-docx, so the tables of contents, tables
and figures are filled with real Word ``PAGEREF`` fields pointing at bookmarks
placed on the headings. Word (or LibreOffice) resolves them when fields are
updated, which is what :func:`litreview.docx_tools.WORD_UPDATE_HINT` tells the
reader to do.
"""
from __future__ import annotations

import copy
from pathlib import Path
from typing import Iterable, Sequence

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

WORD_UPDATE_HINT = ("Page numbers below are Word fields. They were filled in when this document was "
                    "rendered; if you edit the document, select all (Ctrl+A, or Cmd+A on macOS) "
                    "and press F9 to refresh them.")

_BOOKMARK_ID = [1000]


# ----------------------------------------------------------------- inspection

def find_paragraph(doc, text: str, start: int = 0, exact: bool = False):
    """First paragraph whose text contains (or equals) ``text``."""
    t = text.strip().lower()
    for i, p in enumerate(doc.paragraphs[start:], start):
        pt = p.text.strip().lower()
        if (pt == t) if exact else (t in pt):
            return i, p
    raise KeyError(f"paragraph not found: {text!r}")


def paragraph_after(doc, idx: int):
    return doc.paragraphs[idx + 1] if idx + 1 < len(doc.paragraphs) else None


# ----------------------------------------------------------------- text edits

def set_text(par, text: str, *, bold: bool | None = None, size_pt: float | None = None,
             italic: bool | None = None) -> None:
    """Replace a paragraph's text, keeping the first run's formatting."""
    runs = par.runs
    if not runs:
        r = par.add_run(text)
    else:
        runs[0].text = text
        for extra in runs[1:]:
            extra._element.getparent().remove(extra._element)
        r = runs[0]
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    if size_pt is not None:
        r.font.size = Pt(size_pt)


def replace_everywhere(doc, mapping: dict[str, str]) -> dict[str, int]:
    """Replace placeholder strings across paragraphs and tables, run by run.

    Runs are rewritten in place so the template's fonts and sizes survive. A
    placeholder split across several runs is handled by collapsing the paragraph's
    runs when a whole-paragraph match is found.
    """
    counts = {k: 0 for k in mapping}

    def do_par(p):
        for old, new in mapping.items():
            if old in p.text:
                if any(old in (r.text or "") for r in p.runs):
                    for r in p.runs:
                        if old in (r.text or ""):
                            r.text = r.text.replace(old, new)
                            counts[old] += 1
                else:                       # placeholder spans several runs
                    set_text(p, p.text.replace(old, new))
                    counts[old] += 1

    for p in doc.paragraphs:
        do_par(p)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    do_par(p)
    return counts


# --------------------------------------------------------------- construction

def _clone_empty(par):
    """A copy of ``par`` with its runs stripped, so only formatting survives."""
    new = copy.deepcopy(par._element)
    for child in list(new):
        if child.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:bookmarkStart"),
                         qn("w:bookmarkEnd")):
            new.remove(child)
    return new


def insert_after(anchor_par, like_par=None):
    """Insert a new, empty paragraph immediately after ``anchor_par``."""
    from docx.text.paragraph import Paragraph
    el = _clone_empty(like_par or anchor_par)
    anchor_par._element.addnext(el)
    return Paragraph(el, anchor_par._parent)


SUB_OPEN, SUB_CLOSE = "\u27e8", "\u27e9"


def add_runs(par, text: str, *, size_pt: float, bold: bool = False, italic: bool = False):
    """Write ``text`` into ``par``, turning ``x\u27e8c\u27e9`` into a real subscript run.

    Variables appear in the prose as well as in the rendered equations, and
    "n_c" with a literal underscore reads as a typo in a typeset report.
    """
    if SUB_OPEN not in text:
        r = par.add_run(text)
        r.font.size = Pt(size_pt)
        r.bold, r.italic = bold, italic
        return par
    rest = text
    while rest:
        head, sep, tail = rest.partition(SUB_OPEN)
        if head:
            r = par.add_run(head)
            r.font.size = Pt(size_pt)
            r.bold, r.italic = bold, italic
        if not sep:
            break
        subtext, _, rest = tail.partition(SUB_CLOSE)
        r = par.add_run(subtext)
        r.font.size = Pt(size_pt)
        r.font.subscript = True
        r.italic = italic
    return par


def add_body_paragraph(anchor_par, text: str, like_par=None, *, size_pt: float = 12.0,
                       justify: bool = True, first_line_indent: float | None = 0.3,
                       space_after_pt: float = 6.0, bold: bool = False,
                       italic: bool = False, left_indent_in: float | None = None):
    """Body text after ``anchor_par``, inheriting the template's paragraph properties."""
    p = insert_after(anchor_par, like_par)
    pf = p.paragraph_format
    if justify:
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(space_after_pt)
    pf.line_spacing = 1.25
    pf.left_indent = Inches(left_indent_in) if left_indent_in is not None else Inches(0.25)
    pf.first_line_indent = Inches(first_line_indent) if first_line_indent else Inches(0)
    add_runs(p, text, size_pt=size_pt, bold=bold, italic=italic)
    return p


def add_paragraphs(anchor_par, texts: Iterable[str], **kw):
    """Several body paragraphs in order; returns the last one inserted."""
    cur = anchor_par
    for t in texts:
        if not (t or "").strip():
            continue
        cur = add_body_paragraph(cur, t.strip(), **kw)
    return cur


def add_bullets(anchor_par, items: Iterable[str], *, size_pt: float = 12.0,
                numbered: bool = False):
    cur = anchor_par
    for i, item in enumerate(items, 1):
        marker = f"{i}. " if numbered else "• "
        cur = add_body_paragraph(cur, marker + item.strip(), size_pt=size_pt,
                                 first_line_indent=None, left_indent_in=0.5,
                                 space_after_pt=3.0)
    return cur


def add_heading(anchor_par, text: str, *, size_pt: float = 12.0, like_par=None):
    p = insert_after(anchor_par, like_par)
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(10)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.25
    pf.left_indent = Inches(0.25)
    pf.first_line_indent = Inches(0)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size_pt)
    return p


def add_table(doc, anchor_par, rows: Sequence[Sequence[str]], *, header: bool = True,
              size_pt: float = 9.5, style: str = "Table Grid",
              widths_in: Sequence[float] | None = None, align_right_cols: Sequence[int] = ()):
    """A bordered table inserted after ``anchor_par``. ``rows[0]`` is the header."""
    if not rows:
        return None
    n_cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=n_cols)
    try:
        t.style = style
    except KeyError:
        pass
    t.autofit = True
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(n_cols):
            txt = str(row[ci]) if ci < len(row) and row[ci] is not None else ""
            cell = cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            if ci in align_right_cols and ri > 0:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(txt)
            r.font.size = Pt(size_pt)
            r.bold = bool(header and ri == 0)
    if widths_in:
        for ci, w in enumerate(widths_in[:n_cols]):
            for row in t.rows:
                row.cells[ci].width = Inches(w)
    anchor_par._element.addnext(t._element)
    return t


def add_image(anchor_par, image_path: str | Path, *, width_in: float = 5.8):
    """A centred image after ``anchor_par``. Returns the image paragraph."""
    p = insert_after(anchor_par)
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.first_line_indent = Inches(0)
    pf.left_indent = Inches(0)
    pf.space_before = Pt(8)
    pf.space_after = Pt(2)
    p.add_run().add_picture(str(image_path), width=Inches(width_in))
    return p


def add_caption(anchor_par, text: str, *, size_pt: float = 10.0):
    p = insert_after(anchor_par)
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.first_line_indent = Inches(0)
    pf.left_indent = Inches(0)
    pf.space_after = Pt(10)
    r = p.add_run(text)
    r.font.size = Pt(size_pt)
    r.italic = True
    return p


# -------------------------------------------------------- bookmarks and fields

def add_bookmark(par, name: str) -> str:
    """Wrap a paragraph in a bookmark so a PAGEREF field can point at it."""
    _BOOKMARK_ID[0] += 1
    bid = str(_BOOKMARK_ID[0])
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    par._element.insert(0, start)
    par._element.append(end)
    return name


def pageref(par, bookmark: str, placeholder: str = "1") -> None:
    """Append a ``PAGEREF`` field to ``par``; Word fills it in on field update."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f" PAGEREF {bookmark} \\h ")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = placeholder
    run.append(t)
    fld.append(run)
    par._element.append(fld)


def set_cell(cell, text: str, *, bold: bool = False, size_pt: float = 10.0,
             align=None, bookmark_ref: str | None = None) -> None:
    """Write one table cell, optionally ending with a PAGEREF field."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    if align is not None:
        p.paragraph_format.alignment = align
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size_pt)
        r.bold = bold
    if bookmark_ref:
        pageref(p, bookmark_ref)


def ensure_rows(table, n: int):
    """Grow or shrink a template table to exactly ``n`` body rows (header kept)."""
    body = len(table.rows) - 1
    while body < n:
        table.add_row()
        body += 1
    while body > n:
        tr = table.rows[-1]._element
        tr.getparent().remove(tr)
        body -= 1
    return table


def delete_paragraph(par) -> None:
    el = par._element
    el.getparent().remove(el)


def page_break_before(par) -> None:
    pf = par.paragraph_format
    pf.page_break_before = True


# ------------------------------------------------------------------- numbering

class Counter:
    """Figure/table numbering with a register, so captions and lists cannot disagree."""

    def __init__(self, kind: str) -> None:
        self.kind = kind
        self.items: list[dict] = []

    @staticmethod
    def _short(caption: str, limit: int = 120) -> str:
        """First sentence of a caption, for the list of tables / figures."""
        first = caption.split(". ")[0].strip().rstrip(".")
        if len(first) <= limit:
            return first
        cut = first[:limit].rsplit(" ", 1)[0]
        return cut + "\u2026"

    def next(self, caption: str, source: str = "", chapter: int | None = None) -> dict:
        n = len(self.items) + 1
        label = f"{self.kind} {n}"
        item = {"n": n, "label": label, "caption": caption, "short": self._short(caption),
                "source": source, "bookmark": f"_{self.kind.lower()}_{n}", "chapter": chapter}
        self.items.append(item)
        return item

    def __len__(self) -> int:
        return len(self.items)


def set_table_widths(table, widths_in: Sequence[float]) -> None:
    """Pin a table's column widths.

    Word only honours per-cell widths when autofit is off and the table layout is
    fixed, which is why both are set here; the template's tables otherwise
    collapse their merged title column to a few characters wide.
    """
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblLayout", "w:tblW"):
        for el in tbl_pr.findall(qn(tag)):
            tbl_pr.remove(el)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    total = OxmlElement("w:tblW")
    total.set(qn("w:w"), str(int(sum(widths_in) * 1440)))
    total.set(qn("w:type"), "dxa")
    tbl_pr.append(total)

    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        table._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths_in:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(w * 1440)))
        grid.append(col)
    table._tbl.insert(list(table._tbl).index(tbl_pr) + 1, grid)

    for row in table.rows:
        seen = set()
        for cell, w in zip(row.cells, widths_in):
            if id(cell._tc) in seen:        # merged cells repeat in row.cells
                continue
            seen.add(id(cell._tc))
            cell.width = Inches(w)
